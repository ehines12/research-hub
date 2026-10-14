#!/usr/bin/env python3
"""Pull every series for each live dashboard in config/hub.json into data/<dashboard>.json.

Sources: FRED public fredgraph CSV endpoint (no API key; fallback ALFRED), Atlanta Fed GDPNow workbook,
Cleveland Fed inflation-nowcast JSON, TSA checkpoint table, Treasury Fiscal Data DTS API (all keyless).
No numbers are ever synthesized: if a series cannot be fetched, the last good copy on
disk is kept and flagged stale; if none exists, the series is marked failed.

Usage:  python3 fetch_data.py            # all live dashboards
        python3 fetch_data.py macro      # one dashboard
"""
import csv, io, json, sys, time, shutil, subprocess, datetime as dt, urllib.request
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
SOURCES = [
    ("FRED", "https://fred.stlouisfed.org/graph/fredgraph.csv?id={id}"),
    ("ALFRED", "https://alfred.stlouisfed.org/graph/alfredgraph.csv?id={id}"),
]
UA = {"User-Agent": "Mozilla/5.0 (research-hub data refresh)"}


def http_csv(url, tries=3):
    """curl is the primary transport (FRED's edge stalls Python's TLS client); urllib is a fallback."""
    last = None
    for i in range(tries):
        try:
            if shutil.which("curl"):
                r = subprocess.run(["curl", "-sSfL", "--max-time", "60", url], capture_output=True, text=True)
                if r.returncode == 0 and r.stdout:
                    return r.stdout
                raise RuntimeError(f"curl exit {r.returncode}: {r.stderr.strip()[:200]}")
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=60) as r:
                return r.read().decode("utf-8", "replace")
        except Exception as e:  # noqa
            last = e
            time.sleep(2 * (i + 1))
    raise last


def parse(text, sid):
    rdr = csv.reader(io.StringIO(text))
    header = next(rdr)
    if len(header) < 2 or "date" not in header[0].lower():
        raise ValueError(f"unexpected CSV header {header[:3]}")
    obs = []
    for row in rdr:
        if len(row) < 2 or row[1] in ("", "."):
            continue
        try:
            obs.append((row[0], float(row[1])))
        except ValueError:
            continue
    if not obs:
        raise ValueError("no observations")
    return obs


def fetch(sid):
    errs = []
    for name, tmpl in SOURCES:
        try:
            return name, tmpl.format(id=sid), parse(http_csv(tmpl.format(id=sid)), sid)
        except Exception as e:
            errs.append(f"{name}: {e}")
    raise RuntimeError("; ".join(errs))


GDPNOW_PAGE = "https://www.atlantafed.org/cqer/research/gdpnow"
GDPNOW_XLSX = ("https://www.atlantafed.org/-/media/Project/Atlanta/FRBA/Documents/research-and-data/"
               "data/gdpnow/GDPTrackingModelDataAndForecasts.xlsx")


def _qlabel(d):
    return f"Q{(d.month - 1) // 3 + 1} {d.year}"


def fetch_gdpnow():
    """Atlanta Fed GDPNow daily nowcast history from its public tracking workbook.
    Returns (source, url, [(date, value)], {date: {"target": "Q3 2026", "release": "..."}}).
    Finished quarters: TrackingDeepArchives + TrackingArchives ('GDP Nowcast' column).
    Current quarter: CurrentQtrEvolution (Date / Major Releases / GDP* column triplets)."""
    import openpyxl, re, tempfile, os
    fd, path = tempfile.mkstemp(suffix=".xlsx"); os.close(fd)
    try:
        r = subprocess.run(["curl", "-sSfL", "--max-time", "120", "-o", path, GDPNOW_XLSX], capture_output=True, text=True)
        if r.returncode != 0:
            raise RuntimeError(f"curl exit {r.returncode}: {r.stderr.strip()[:200]}")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        pts = {}  # date -> (value, target_quarter_end, release)
        for name in ("TrackingDeepArchives", "TrackingArchives"):
            rows = wb[name].iter_rows(values_only=True)
            hdr = next(rows)
            gi = [i for i, c in enumerate(hdr) if c and str(c).strip().lower() == "gdp nowcast"]
            if not gi:
                raise ValueError(f"{name}: no 'GDP Nowcast' column")
            for row in rows:
                d, q, v = row[0], row[1], row[gi[0]]
                if hasattr(d, "year") and hasattr(q, "year") and isinstance(v, (int, float)):
                    if d.date() not in pts or q.date() > pts[d.date()][1]:
                        pts[d.date()] = (float(v), q.date(), None)
        title = next(wb["TrackingHistory"].iter_rows(min_row=2, max_row=2, values_only=True))[0] or ""
        m = re.search(r"(\d{4})\s*q\s*(\d)", str(title), re.I)
        if not m:
            raise ValueError(f"cannot find current quarter in TrackingHistory title {title!r}")
        y, qn = int(m.group(1)), int(m.group(2))
        cur_q = dt.date(y, 3 * qn, 31 if qn in (1, 4) else 30)
        rows = list(wb["CurrentQtrEvolution"].iter_rows(values_only=True))
        if not rows or "date" not in str(rows[0][0]).lower():
            raise ValueError("CurrentQtrEvolution layout changed")
        n_cur = 0
        for row in rows[1:]:
            for b in range(0, len(row) - 2, 3):
                d, lab, v = row[b], row[b + 1], row[b + 2]
                if hasattr(d, "year") and isinstance(v, (int, float)):
                    pts[d.date()] = (float(v), cur_q, str(lab).strip() if lab else None); n_cur += 1
        if n_cur == 0:
            raise ValueError("no current-quarter nowcasts found")
    finally:
        os.unlink(path)
    dates = sorted(pts)
    obs = [(d.isoformat(), pts[d][0]) for d in dates]
    meta = {d.isoformat(): {"target": _qlabel(pts[d][1]), **({"release": pts[d][2]} if pts[d][2] else {})} for d in dates}
    return "Atlanta Fed", GDPNOW_XLSX, obs, meta


# ---------------------------------------------------------------------------------------------
# Non-FRED sources. Each fetcher returns (source, url, [(date, value)], meta|None). Downloads that
# feed several series (one Cleveland Fed file -> 4 measures, one TSA pull -> level + y/y) are cached
# per run so each URL is fetched once.
import threading, re as _re
_CACHE, _LOCKS, _GL = {}, {}, threading.Lock()


def _cached(key, fn):
    with _GL:
        lk = _LOCKS.setdefault(key, threading.Lock())
    with lk:
        if key not in _CACHE:
            try:
                _CACHE[key] = ("ok", fn())
            except Exception as e:  # noqa: BLE001 - re-raised for every dependent series
                _CACHE[key] = ("err", e)
    st, v = _CACHE[key]
    if st == "err":
        raise v
    return v


def _curl(url, timeout=120, ua=None):
    cmd = ["curl", "-sSfL", "-g", "--max-time", str(timeout), "--retry", "2", "--retry-delay", "3"]
    if ua:
        cmd += ["-A", ua]
    r = subprocess.run(cmd + [url], capture_output=True)
    if r.returncode != 0 or not r.stdout:
        raise RuntimeError(f"curl exit {r.returncode}: {r.stderr.decode(errors='replace').strip()[:200]}")
    return r.stdout


# Cleveland Fed inflation nowcast: the page's charts read these public JSON files (FusionCharts
# configs, one per target period, each with daily nowcasts and the actual print once released).
CLEV_PAGE = "https://www.clevelandfed.org/indicators-and-data/inflation-nowcasting"
CLEV_JSON = "https://www.clevelandfed.org/-/media/files/webcharts/inflationnowcasting/nowcast_{h}.json"
MON = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()


def _clev_targets(h):
    """{target_label: {"start": date, "measures": {name: {"pts": [(date, v)], "actual": v|None}}}}, as_of"""
    doc = json.loads(_curl(CLEV_JSON.format(h=h)))
    out, as_of = {}, None
    for e in doc:
        sub = e["chart"]["subcaption"].strip()
        if h == "quarter":
            m = _re.fullmatch(r"(\d{4}):Q([1-4])", sub); y, mo = int(m.group(1)), 3 * int(m.group(2)) - 2
            label = f"Q{m.group(2)} {y}"
        else:
            m = _re.fullmatch(r"(\d{4})-(\d{1,2})", sub); y, mo = int(m.group(1)), int(m.group(2))
            label = f"{MON[mo - 1]} {y}"
        if not m:
            raise ValueError(f"unexpected target {sub!r}")
        # category labels are mm/dd without a year: infer it from the target period, rolling over at Jan
        days, yy, pm = [], y, mo
        for c in e["categories"][0]["category"]:
            if c.get("vline"):
                continue  # release markers ("CPI Aug"), no data point
            mm, dd = map(int, c["label"].split("/"))
            if mm < pm:
                yy += 1
            pm = mm
            days.append(dt.date(yy, mm, dd))
        if any(b <= a for a, b in zip(days, days[1:])) or (days and (days[-1] - days[0]).days > 220):
            raise ValueError(f"{sub}: non-monotonic dates")
        meas = {}
        for ds in e["dataset"]:
            name = ds["seriesname"].strip()
            if len(ds["data"]) != len(days):
                raise ValueError(f"{sub}/{name}: {len(ds['data'])} points vs {len(days)} dates")
            pts = [(days[i], float(x["value"])) for i, x in enumerate(ds["data"]) if str(x.get("value", "")).strip() not in ("",)]
            if name.startswith("Actual "):
                meas.setdefault(name[7:], {"pts": [], "actual": None})["actual"] = pts[-1][1] if pts else None
            else:
                meas.setdefault(name, {"pts": [], "actual": None})["pts"] = pts
                if pts and (as_of is None or pts[-1][0] > as_of):
                    as_of = pts[-1][0]
        out[label] = {"start": dt.date(y, mo, 1), "measures": meas}
    if not out or as_of is None:
        raise ValueError("no nowcasts found")
    return out, as_of


def fetch_clevnow(s):
    """mode 'nowcast': one obs per target period = latest nowcast (final pre-release one for past targets).
       mode 'actual': the released print the Cleveland Fed charts it against."""
    h, name, mode = s["horizon"], s["measure"], s.get("mode", "nowcast")
    tg, as_of = _cached("clev:" + h, lambda: _clev_targets(h))
    mm = _cached("clev:month", lambda: _clev_targets("month"))[0] if s.get("with_mm") else None
    obs, meta = [], {}
    for label, t in sorted(tg.items(), key=lambda kv: kv[1]["start"]):
        m = t["measures"].get(name)
        if not m:
            continue
        k = t["start"].isoformat()
        if mode == "actual":
            if m["actual"] is not None:
                obs.append((k, m["actual"]))
            continue
        if not m["pts"]:
            continue
        last_d, last_v = m["pts"][-1]
        obs.append((k, last_v))
        pend = m["actual"] is None and last_d == as_of
        md = {"target": label, "asof": last_d.isoformat(), "pending": pend}
        wk = [p for p in m["pts"] if p[0] <= last_d - dt.timedelta(days=7)]
        if pend and wk:
            md.update(prior_date=wk[-1][0].isoformat(), prior_value=round(wk[-1][1], 4))
        if m["actual"] is not None:
            md["actual"] = round(m["actual"], 4)
        if mm and label in mm and mm[label]["measures"].get(name, {}).get("pts"):
            md["mm"] = round(mm[label]["measures"][name]["pts"][-1][1], 4)
        meta[k] = md
    if not obs:
        raise ValueError(f"no {mode} values for {name}")
    return "Cleveland Fed", CLEV_PAGE, obs, meta


# TSA checkpoint travel numbers: the published HTML table (current year page + one page per past year).
TSA_URL = "https://www.tsa.gov/travel/passenger-volumes"
BROWSER_UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/140.0 Safari/537.36"


def _tsa_daily():
    this = dt.date.today().year
    pts = {}
    for url in [TSA_URL] + [f"{TSA_URL}/{y}" for y in (this - 1, this - 2, this - 3)]:
        html = _curl(url, 60, BROWSER_UA).decode("utf-8", "replace")
        cells = _re.findall(r"<td[^>]*>\s*([^<]*?)\s*</td>", html)
        n = 0
        for a, b in zip(cells[::2], cells[1::2]):
            m = _re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", a)
            if m and _re.fullmatch(r"[\d,]+", b):
                pts[dt.date(int(m.group(3)), int(m.group(1)), int(m.group(2)))] = int(b.replace(",", "")); n += 1
        if n < 25:
            raise ValueError(f"TSA table layout changed ({url}: {n} rows)")
    return pts


def fetch_tsa(s):
    """mode 'avg7': trailing 7-day average (millions); 'yoy': that average vs the same 7 days 364 days
       earlier (same weekdays), %. Only full 7-day windows are used."""
    pts = _cached("tsa", _tsa_daily)
    days = sorted(pts)
    avg = {}
    for d0 in days:
        w = [pts.get(d0 - dt.timedelta(days=i)) for i in range(7)]
        if all(v is not None for v in w):
            avg[d0] = sum(w) / 7.0
    if s.get("mode") == "yoy":
        obs = [(d0.isoformat(), (v / avg[d0 - dt.timedelta(days=364)] - 1) * 100) for d0, v in sorted(avg.items()) if d0 - dt.timedelta(days=364) in avg]
    else:
        obs = [(d0.isoformat(), v / 1e6) for d0, v in sorted(avg.items())]
    if not obs:
        raise ValueError("TSA: not enough history")
    return "TSA", TSA_URL, obs, None


# Daily Treasury Statement, Table II deposits: 'Taxes - Withheld Individual/FICA' ($ millions).
DTS_API = ("https://api.fiscaldata.treasury.gov/services/api/fiscal_service/v1/accounting/dts/"
           "deposits_withdrawals_operating_cash?fields=record_date,transaction_catg,transaction_today_amt"
           "&filter=transaction_catg:eq:Taxes%20-%20Withheld%20Individual/FICA,record_date:gte:2023-01-01"
           "&sort=record_date&page[size]=10000")


def _dts_daily():
    doc = json.loads(_curl(DTS_API, 90))
    rows = doc.get("data") or []
    if len(rows) < 260:
        raise ValueError(f"DTS: only {len(rows)} rows")
    pts = {}
    for r in rows:
        pts[d(r["record_date"])] = pts.get(d(r["record_date"]), 0.0) + float(r["transaction_today_amt"])
    return pts


def fetch_dts(s):
    """mode 'sum20': rolling N-business-day (N DTS reports, default 20) sum, $bn; 'yoy': that sum vs the sum
       ending on the last report on/before the same date a year earlier, %."""
    pts = _cached("dts", _dts_daily)
    days = sorted(pts)
    n = int(s.get("window", 20))
    roll = {days[i]: sum(pts[x] for x in days[i - n + 1:i + 1]) / 1000.0 for i in range(n - 1, len(days))}
    if s.get("mode") == "yoy":
        rd = sorted(roll)
        obs = []
        for d0 in rd:
            try:
                tgt = d0.replace(year=d0.year - 1)
            except ValueError:
                tgt = d0 - dt.timedelta(days=365)
            prev = [x for x in rd if x <= tgt and (tgt - x).days <= 5]
            if prev:
                obs.append((d0.isoformat(), (roll[d0] / roll[prev[-1]] - 1) * 100))
    else:
        obs = [(x.isoformat(), v) for x, v in sorted(roll.items())]
    if not obs:
        raise ValueError("DTS: not enough history")
    return "Treasury Fiscal Data (DTS)", "https://fiscaldata.treasury.gov/datasets/daily-treasury-statement/", obs, None



# Caldara-Iacoviello Geopolitical Risk (GPR) and Trade Policy Uncertainty (TPU) indexes (public xls/xlsx).
GPR_DAILY = "https://www.matteoiacoviello.com/gpr_files/data_gpr_daily_recent.xls"
GPR_MONTHLY = "https://www.matteoiacoviello.com/gpr_files/data_gpr_export.xls"
TPU_XLSX = "https://www.matteoiacoviello.com/tpu_files/tpu_web_latest.xlsx"


def _xls_table(url):
    import xlrd
    b = xlrd.open_workbook(file_contents=_curl(url, 120, BROWSER_UA))
    sh = b.sheet_by_index(0)
    hdr = [str(h).strip() for h in sh.row_values(0)]
    return hdr, [sh.row_values(i) for i in range(1, sh.nrows)], b.datemode


def fetch_gpr(s):
    """s['file'] = daily|monthly, s['col'] = GPRD / GPRD_ACT / GPRD_THREAT (daily) or GPR / GPRT / GPRA (monthly)."""
    import xlrd
    url = GPR_DAILY if s["file"] == "daily" else GPR_MONTHLY
    hdr, rows, mode = _cached("gpr:" + s["file"], lambda: _xls_table(url))
    if s["col"] not in hdr:
        raise ValueError(f"GPR {s['file']}: column {s['col']} missing (have {hdr[:12]})")
    ci, obs = hdr.index(s["col"]), []
    for r in rows:
        v = r[ci]
        if v in ("", None):
            continue
        if s["file"] == "daily":
            t = str(r[hdr.index("DAY")]).split(".")[0]
            if not _re.fullmatch(r"\d{8}", t):
                continue
            d0 = f"{t[:4]}-{t[4:6]}-{t[6:]}"
        else:
            m = r[hdr.index("month")]
            if not isinstance(m, float):
                continue
            d0 = xlrd.xldate_as_datetime(m, mode).date().replace(day=1).isoformat()
        obs.append((d0, float(v)))
    if len(obs) < 24:
        raise ValueError("GPR: too few observations")
    return "Caldara-Iacoviello GPR", "https://www.matteoiacoviello.com/gpr.htm", sorted(obs), None


def fetch_tpu(s):
    """Caldara et al. Trade Policy Uncertainty: s['sheet'] = TPU_DAILY (col TPUD_index) or TPU_MONTHLY (col TPU)."""
    import openpyxl, io as _io
    wb = _cached("tpu", lambda: openpyxl.load_workbook(_io.BytesIO(_curl(TPU_XLSX, 120, BROWSER_UA)), read_only=True, data_only=True))
    with _GL:
        rows = list(wb[s["sheet"]].iter_rows(values_only=True))
    hdr = [str(h).strip() if h else "" for h in rows[0]]
    if s["col"] not in hdr:
        raise ValueError(f"TPU {s['sheet']}: column {s['col']} missing")
    ci, obs = hdr.index(s["col"]), []
    for r in rows[1:]:
        if hasattr(r[0], "year") and isinstance(r[ci], (int, float)):
            obs.append((r[0].date().isoformat(), float(r[ci])))
    if len(obs) < 24:
        raise ValueError("TPU: too few observations")
    return "Caldara et al. TPU", "https://www.matteoiacoviello.com/tpu.htm", sorted(obs), None


# IMF PortWatch daily chokepoint transits (public ArcGIS feature service; ~1 week publication lag).
PORTWATCH = ("https://services9.arcgis.com/weJ1QsnbMYJlCHdG/arcgis/rest/services/Daily_Chokepoints_Data/"
             "FeatureServer/0/query")


def _portwatch_rows(portid):
    rows, off = [], 0
    while True:
        url = (f"{PORTWATCH}?where=portid%3D%27{portid}%27&outFields=date,portname,n_total,n_tanker,n_container,n_dry_bulk,capacity"
               f"&orderByFields=date&resultOffset={off}&resultRecordCount=1000&f=json")
        for attempt in range(4):
            d = json.loads(_curl(url, 90))
            err = d.get("error")
            if err and err.get("code") == 429 and attempt < 3:
                time.sleep(65)  # ArcGIS per-minute quota; wait it out and retry
                continue
            break
        if d.get("error"):
            raise RuntimeError(f"PortWatch: {d['error']}")
        time.sleep(2)  # pace page requests to stay under the per-minute quota
        f = [x["attributes"] for x in d.get("features", [])]
        rows += f
        if len(f) < 1000 and not d.get("exceededTransferLimit"):
            break
        off += len(f)
    if len(rows) < 400:
        raise ValueError(f"PortWatch {portid}: only {len(rows)} rows")
    return rows


def fetch_portwatch(s):
    rows = _cached("pw:" + s["portid"], lambda: _portwatch_rows(s["portid"]))
    obs = sorted({str(r["date"])[:10]: float(r[s["field"]]) for r in rows if r.get(s["field"]) is not None}.items())
    return "IMF PortWatch", "https://portwatch.imf.org/pages/chokepoints", obs, None


def _fetch_series(s):
    src = s.get("source")
    if src == "gdpnow":
        return fetch_gdpnow()
    if src == "clevnow":
        return fetch_clevnow(s)
    if src == "tsa":
        return fetch_tsa(s)
    if src == "dts":
        return fetch_dts(s)
    if src == "gpr":
        return fetch_gpr(s)
    if src == "tpu":
        return fetch_tpu(s)
    if src == "portwatch":
        return fetch_portwatch(s)
    return fetch(s["fred"]) + (None,)


def _safe_fetch(spec):
    try:
        return _fetch_series(spec)
    except Exception as e:
        return e


def d(s):
    return dt.date.fromisoformat(s)


def transform(obs, how):
    if how == "level":
        return obs
    if how == "div1000":
        return [(t, v / 1000.0) for t, v in obs]
    if how == "diff":
        return [(obs[i][0], obs[i][1] - obs[i - 1][1]) for i in range(1, len(obs))]
    if how == "yoy":
        idx = {t: v for t, v in obs}
        out = []
        for t, v in obs:
            dd = d(t)
            try:
                prev = idx.get(dd.replace(year=dd.year - 1).isoformat())
            except ValueError:
                prev = None
            if prev:
                out.append((t, (v / prev - 1) * 100.0))
        return out
    if how == "mavg":  # weekly/daily -> calendar-month average, dated the 1st
        acc = {}
        for t, v in obs:
            acc.setdefault(t[:7] + "-01", []).append(v)
        return [(k, round(sum(v) / len(v), 4)) for k, v in sorted(acc.items())]
    m = _re.fullmatch(r"ma(\d+)(_ly)?", how)
    if m:
        n = int(m.group(1))
        out = [(obs[i][0], sum(v for _, v in obs[i - n + 1:i + 1]) / n) for i in range(n - 1, len(obs))
               if (d(obs[i][0]) - d(obs[i - n + 1][0])).days <= n + 2]  # skip windows spanning data gaps
        if m.group(2):  # shift forward 364 days (same weekday) to overlay last year's path
            out = [((d(t) + dt.timedelta(days=364)).isoformat(), v) for t, v in out]
            out = [(t, v) for t, v in out if t <= obs[-1][0]]
        return out
    raise ValueError(f"unknown transform {how}")


def year_ago(obs):
    """Observation closest to (and not after) one year before the latest."""
    target = d(obs[-1][0]).replace(year=d(obs[-1][0]).year - 1) if not (d(obs[-1][0]).month == 2 and d(obs[-1][0]).day == 29) else d(obs[-1][0]) - dt.timedelta(days=365)
    cand = [o for o in obs if d(o[0]) <= target]
    return cand[-1] if cand else None


def run_predictions(cfg, out_path, prev):
    """Prediction-markets dashboard (predictions.py). If either venue fails, keep the last good list
    (flagged stale) rather than ranking on one venue only."""
    import predictions
    now = dt.datetime.now(dt.timezone.utc)
    r = predictions.build(cfg)
    src = r["sources"]
    bad = [k for k, v in src.items() if v.get("status") != "ok"]
    if bad and prev.get("items"):
        doc = {**prev, "config": cfg, "stale": True, "sources": src}
    else:
        doc = {"id": cfg["id"], "config": cfg, "updated_at": now.isoformat(timespec="seconds"), **r}
    doc["series"] = {k: {"fred": k.capitalize(), "status": ("stale" if (bad and prev.get("items")) else "failed") if v.get("status") != "ok" else "ok",
                         **({"error": v["error"]} if v.get("error") else {})} for k, v in src.items()}
    for it in doc.get("items", []):
        print(f"  {it['rank']:>2} {it['group']:<12} {it['vol7d']/1e3:>8.0f}k  {it['title'][:70]}")
    print("  sources:", json.dumps(src))
    DATA.mkdir(exist_ok=True)
    out_path.write_text(json.dumps(doc, separators=(",", ":")))
    return doc


def run_dashboard(cfg_path):
    cfg = json.loads((ROOT / cfg_path).read_text())
    if cfg.get("kind") == "bank_conditions":
        import bank_conditions
        doc = bank_conditions.build(cfg)
        c = doc["current"]
        print(f"  diffusion {c['score']} ({c['label']}), {c['n']}/{c['total']} scored")
        out = DATA / f"{cfg['id']}.json"
        # Guard: if a source outage (e.g. FRED 403s) leaves under 75% of inputs scored, keep the last good file
        # instead of publishing an index built from a handful of inputs.
        if c["total"] and c["n"] < 0.75 * c["total"] and out.exists():
            print(f"  -> only {c['n']}/{c['total']} inputs scored (source outage); kept previous {out.name}")
            return json.loads(out.read_text())
        out.write_text(json.dumps(doc, separators=(",", ":")))
        return doc
    if cfg.get("kind") == "predictions":
        out_path = DATA / f"{cfg['id']}.json"
        prev = json.loads(out_path.read_text()) if out_path.exists() else {}
        return run_predictions(cfg, out_path, prev)
    out_path = DATA / f"{cfg['id']}.json"
    prev = json.loads(out_path.read_text()) if out_path.exists() else {"series": {}}
    now = dt.datetime.now(dt.timezone.utc)
    result = {"id": cfg["id"], "config": cfg, "updated_at": now.isoformat(timespec="seconds"), "series": {}}
    with ThreadPoolExecutor(max_workers=6) as ex:
        fetched = dict(zip([s["key"] for s in cfg["series"]],
                           ex.map(_safe_fetch, cfg["series"])))
    for s in cfg["series"]:
        rec = {k: v for k, v in s.items()}
        try:
            res = fetched[s["key"]]
            if isinstance(res, Exception):
                raise res
            src, url, raw, meta = res
            obs = [(t, round(v, 4)) for t, v in transform(raw, s["transform"])]
            if meta:
                rec["obs_meta"] = {t: meta[t] for t, _ in obs[-400:] if t in meta}
            if len(raw) < 2:
                raise ValueError("fewer than 2 observations")
            rec.update(status="ok", source=src, source_url=url, fetched_at=now.isoformat(timespec="seconds"),
                       raw_latest={"date": raw[-1][0], "value": raw[-1][1]}, obs=obs)
        except Exception as e:
            old = prev.get("series", {}).get(s["key"])
            if old and old.get("obs"):
                rec.update(old)
                rec.update(status="stale", error=str(e))
            else:
                rec.update(status="failed", error=str(e), obs=[])
        if rec["obs"]:
            o = rec["obs"]
            rec["latest"] = {"date": o[-1][0], "value": o[-1][1]}
            rec["prior"] = {"date": o[-2][0], "value": o[-2][1]} if len(o) > 1 else None
            ya = year_ago(o)
            rec["year_ago"] = {"date": ya[0], "value": ya[1]} if ya and not s.get("no_yoy") else None
            om = rec.get("obs_meta")
            if om and o[-1][0] in om:
                # nowcast-style series: prior = previous estimate for the SAME target quarter
                m0 = om[o[-1][0]]
                rec["latest"].update(m0)
                if "asof" in m0:  # one obs per target period (Cleveland Fed): prior = estimate ~1 week earlier
                    rec["prior"] = {"date": m0["prior_date"], "value": m0["prior_value"]} if "prior_value" in m0 else None
                else:  # daily nowcast (GDPNow): prior = previous estimate for the SAME target quarter
                    tgt = m0.get("target")
                    same = [x for x in o[:-1] if om.get(x[0], {}).get("target") == tgt]
                    rec["prior"] = {"date": same[-1][0], "value": same[-1][1], **om[same[-1][0]]} if same else None
        result["series"][s["key"]] = rec
        st = rec["status"]
        lt = rec.get("latest") or {}
        print(f"  {s.get('fred') or s['key']:<16} {st:<7} {lt.get('date','-'):<11} {lt.get('value','')}  {rec.get('error','')}")
    for s in cfg["series"]:  # card statistics + optional obs trimming (raw daily series used only for stats)
        rec = result["series"][s["key"]]
        if s.get("stats") and rec.get("obs"):
            try:
                rec["stats"] = _stats(s["stats"], rec["obs"], s.get("freq"))
            except Exception as e:  # noqa: BLE001
                rec["stats_error"] = str(e)
        if s.get("keep_obs") and rec.get("obs"):
            rec["obs"] = rec["obs"][-int(s["keep_obs"]):]
    if cfg.get("auctions"):
        _auctions_block(cfg["auctions"], result, prev)
    _derive_all(cfg, result)
    if cfg.get("deposit_pricing"):
        try:
            result["deposit_pricing"] = _deposit_block(cfg["deposit_pricing"], result["series"])
        except Exception as e:  # noqa: BLE001
            result["deposit_pricing"] = {"error": str(e)}
            print(f"  deposit_pricing FAILED: {e}")
    DATA.mkdir(exist_ok=True)
    # Only rewrite the (large) data file when the numbers or config actually changed, so the
    # scheduled job commits a big diff only on new FRED data. updated_at = last data change.
    changed = _strip(result) != _strip(prev)
    if result.get("quarterly") and (changed or not (DATA / f"{cfg['id']}-quarterly.csv").exists()):
        _write_quarterly_files(cfg["id"], result["quarterly"])
    if result.get("growth5y") and (changed or not (DATA / f"{cfg['id']}-growth5y.csv").exists()):
        _write_growth5y_files(cfg["id"], result["growth5y"])
    if changed:
        tmp = out_path.with_suffix(".tmp")
        tmp.write_text(json.dumps(result, separators=(",", ":")))
        tmp.replace(out_path)
    else:
        result["updated_at"] = prev.get("updated_at", result["updated_at"])
    print(f"  -> {'data changed, wrote' if changed else 'no data change, kept'} {out_path.name}")
    return result


def _stats(kind, obs, freq):
    """'risk': latest, short average (30 obs daily / 3 months monthly) vs trailing 1Y average, and the percentile of
    that short average among all its historical values. 'chokepoint' (on a 7-day-average series): latest vs the same
    date a year earlier (364 days, same weekday) and vs the 2019-2022 average for the same calendar date."""
    vals = [v for _, v in obs]
    last_d = d(obs[-1][0])
    if kind == "risk":
        n = 30 if freq == "Daily" else 3
        roll = [sum(vals[i - n + 1:i + 1]) / n for i in range(n - 1, len(vals))]
        cut = last_d - dt.timedelta(days=365)
        yr = [v for t, v in obs if d(t) > cut]
        cur = roll[-1]
        pct = 100.0 * sum(1 for x in roll if x <= cur) / len(roll)
        return {"latest": vals[-1], "date": obs[-1][0], "short": round(cur, 2), "short_n": n, "avg1y": round(sum(yr) / len(yr), 2),
                "pctile": round(pct, 1), "since": obs[0][0][:4]}
    if kind == "chokepoint":
        idx = {t: v for t, v in obs}
        def near(day):
            for k in range(4):
                for dd in (day - dt.timedelta(days=k), day + dt.timedelta(days=k)):
                    if dd.isoformat() in idx:
                        return idx[dd.isoformat()]
            return None
        ly = near(last_d - dt.timedelta(days=364))
        base = [near(last_d.replace(year=y) if not (last_d.month == 2 and last_d.day == 29) else last_d.replace(year=y, day=28)) for y in (2019, 2020, 2021, 2022)]
        base = [b for b in base if b is not None]
        b0 = sum(base) / len(base) if base else None
        return {"latest": vals[-1], "date": obs[-1][0], "ly": ly, "base": round(b0, 2) if b0 is not None else None,
                "vs_ly": round((vals[-1] / ly - 1) * 100, 1) if ly else None, "vs_base": round((vals[-1] / b0 - 1) * 100, 1) if b0 else None}
    raise ValueError(f"unknown stats kind {kind}")


# ---- Treasury coupon auctions: Fiscal Data auctions_query (results) + TreasuryDirect (upcoming) ----
AUCTIONS_API = ("https://api.fiscaldata.treasury.gov/services/api/fiscal_service/v1/accounting/od/auctions_query"
                "?filter=auction_date:gte:{since},security_type:in:(Note,Bond)&sort=-auction_date&page[size]=10000")
TD_UPCOMING = "https://www.treasurydirect.gov/TA_WS/securities/upcoming?format=json"


def _num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def _tenor(term, tips, frn):
    t = (term or "").strip()
    m = _re.match(r"(\d+)-Year", t)
    y = m.group(1) + "Y" if m else t
    return y + (" TIPS" if tips == "Yes" else "") + (" FRN" if frn == "Yes" else "")


def _auctions_block(ac, result, prev):
    key = "TREASURY_AUCTIONS"
    now = dt.datetime.now(dt.timezone.utc)
    try:
        since = (dt.date.today() - dt.timedelta(days=int(ac.get("history_days", 1100)))).isoformat()
        rows = json.loads(_curl(AUCTIONS_API.format(since=since), 90)).get("data") or []
        if len(rows) < 50:
            raise ValueError(f"only {len(rows)} auction rows")
        recs = []
        for r in rows:
            if _num(r.get("bid_to_cover_ratio")) is None:
                continue  # announced, not yet auctioned
            frn = r.get("floating_rate") == "Yes"
            comp = _num(r.get("comp_accepted")) or 0
            pctof = lambda k: round(100 * (_num(r.get(k)) or 0) / comp, 1) if comp else None
            recs.append({"date": r["auction_date"], "cusip": r["cusip"], "type": r["security_type"],
                         "term": r["security_term"], "tenor": _tenor(r.get("original_security_term") if r.get("original_security_term") not in (None, "", "null") else r["security_term"], r.get("inflation_index_security"), r.get("floating_rate")),
                         "reopening": r.get("reopening") == "Yes", "size_bn": round((_num(r.get("offering_amt")) or 0) / 1e9, 1),
                         "yield": _num(r.get("high_discnt_margin")) if frn else _num(r.get("high_yield")), "yield_kind": "disc. margin" if frn else ("real yield" if r.get("inflation_index_security") == "Yes" else "yield"),
                         "btc": _num(r.get("bid_to_cover_ratio")), "indirect": pctof("indirect_bidder_accepted"), "direct": pctof("direct_bidder_accepted"),
                         "dealer": pctof("primary_dealer_accepted")})
        recs.sort(key=lambda x: x["date"])
        n = int(ac.get("compare_n", 6))
        for i, x in enumerate(recs):
            prior = [y for y in recs[:i] if y["tenor"] == x["tenor"]][-n:]
            if len(prior) >= 3:
                av = lambda k: round(sum(y[k] for y in prior if y[k] is not None) / len([y for y in prior if y[k] is not None]), 2)
                x["avg"] = {"btc": av("btc"), "indirect": av("indirect"), "direct": av("direct"), "dealer": av("dealer"), "n": len(prior)}
                db, dd = x["btc"] - x["avg"]["btc"], x["dealer"] - x["avg"]["dealer"]
                x["d_btc"], x["d_dealer"], x["d_indirect"] = round(db, 2), round(dd, 1), round(x["indirect"] - x["avg"]["indirect"], 1)
                weak, strong = db <= -0.10 or dd >= 3.0, db >= 0.10 or dd <= -3.0
                x["verdict"] = "weak" if weak and not (db >= 0.10 and dd <= -3.0) else ("strong" if strong else "in line")
        tenors = ac.get("tenors")
        recent = [x for x in reversed(recs) if not tenors or x["tenor"] in tenors][: int(ac.get("recent_n", 24))]
        up = []
        try:
            for u in json.loads(_curl(TD_UPCOMING, 60)):
                if u.get("securityType") not in ("Note", "Bond"):
                    continue
                up.append({"date": u["auctionDate"][:10], "issue": (u.get("issueDate") or "")[:10], "type": u["securityType"], "term": u.get("securityTerm"),
                           "tenor": _tenor(u.get("originalSecurityTerm") or u.get("securityTerm"), u.get("tips"), u.get("floatingRate")),
                           "reopening": u.get("reopening") == "Yes", "size_bn": round(_num(u.get("offeringAmount")) / 1e9, 1) if _num(u.get("offeringAmount")) else None,
                           "announced": (u.get("announcementDate") or "")[:10]})
            up_status = "ok"
        except Exception as e:  # noqa: BLE001
            up_status = f"failed: {e}"
        result["auctions"] = {"recent": recent, "upcoming": sorted(up, key=lambda x: x["date"]), "upcoming_status": up_status,
                              "note": ac.get("note"), "latest": recent[0]["date"] if recent else None}
        st = "ok" if up_status == "ok" else "stale"
        result["series"][key] = {"key": key, "fred": "Treasury auctions", "status": st, "title": "Treasury auctions", "obs": [],
                                 "latest": {"date": recent[0]["date"], "value": recent[0]["btc"]} if recent else None, **({"error": up_status} if st != "ok" else {})}
    except Exception as e:  # noqa: BLE001
        old = prev.get("auctions")
        if old:
            result["auctions"] = {**old, "stale": True}
        result["series"][key] = {"key": key, "fred": "Treasury auctions", "status": "stale" if old else "failed", "title": "Treasury auctions", "obs": [], "error": str(e)[:300]}
    lt = (result["series"][key].get("latest") or {})
    print(f"  {'TREASURY_AUCTIONS':<16} {result['series'][key]['status']:<7} {lt.get('date','-'):<11} {len((result.get('auctions') or {}).get('recent', []))} recent, {len((result.get('auctions') or {}).get('upcoming', []))} upcoming")


def _cycle_points(ff, start="2021-06-01", thr=0.20):
    """Turning points in monthly fed funds since `start`: alternating troughs/peaks, each confirmed once the rate
    reverses by >= thr pp. The point is the LAST month at the extreme level (i.e. the month before the first move)."""
    ff = [(t, v) for t, v in ff if t >= start]
    pts, mode, ext_i = [], "trough", 0
    for i, (t, v) in enumerate(ff):
        ev = ff[ext_i][1]
        if mode == "trough":
            if v <= ev:
                ext_i = i
            elif v - ev >= thr:
                pts.append(("trough", ff[ext_i][0], ev)); mode, ext_i = "peak", i
        else:
            if v >= ev:
                ext_i = i
            elif ev - v >= thr:
                pts.append(("peak", ff[ext_i][0], ev)); mode, ext_i = "trough", i
    return pts


def _deposit_block(dp, S):
    ffk = dp["fedfunds"]
    ff = [tuple(o) for o in S[ffk]["obs"]]
    ffm = dict(ff)
    prods = [(p["key"], p["label"]) for p in dp["products"]]
    rates = {k: dict(tuple(o) for o in S[k].get("obs") or []) for k, _ in prods}
    last = min(max(r) for r in rates.values() if r)  # latest month every product has
    pts = _cycle_points(ff)
    legs = []
    for a, b in zip(pts, pts[1:] + [("latest", last, ffm.get(last))]):
        if b[1] <= a[1] or ffm.get(b[1]) is None:
            continue
        kind = "Hiking" if a[0] == "trough" else "Easing"
        legs.append({"kind": kind, "start": a[1], "end": b[1], "end_is_latest": b[0] == "latest",
                     "ff_start": ffm[a[1]], "ff_end": ffm[b[1]], "d_ff": round(ffm[b[1]] - ffm[a[1]], 2)})
    if len(pts) >= 1:  # full cycle: first trough -> latest
        a = pts[0]
        legs.append({"kind": "Full cycle", "start": a[1], "end": last, "end_is_latest": True, "ff_start": ffm[a[1]], "ff_end": ffm[last], "d_ff": round(ffm[last] - ffm[a[1]], 2)})
    for L in legs:
        L["betas"] = {}
        for k, _ in prods:
            r0, r1 = rates[k].get(L["start"]), rates[k].get(L["end"])
            if r0 is not None and r1 is not None and abs(L["d_ff"]) >= 0.25:
                L["betas"][k] = {"beta": round((r1 - r0) / L["d_ff"] * 100, 1), "d_rate": round(r1 - r0, 2), "r0": round(r0, 2), "r1": round(r1, 2)}
    qs = sorted({t for t in rates[prods[0][0]] if t[5:7] in ("03", "06", "09", "12") and t <= last})[-8:]
    table = {"quarters": qs, "rows": [{"key": k, "label": lab, "vals": [rates[k].get(q) for q in qs]} for k, lab in prods] +
             [{"key": ffk, "label": "Effective fed funds (monthly avg)", "vals": [ffm.get(q) for q in qs]}]}
    print(f"  deposit_pricing: {len(legs)} legs {[ (l['kind'], l['start'][:7], l['end'][:7]) for l in legs]}, latest {last}")
    return {"latest": last, "turning_points": pts, "legs": legs, "table": table, "note": dp.get("note")}


def _latest_fields(rec, no_yoy=False):
    o = rec["obs"]
    rec["latest"] = {"date": o[-1][0], "value": o[-1][1]}
    rec["prior"] = {"date": o[-2][0], "value": o[-2][1]} if len(o) > 1 else None
    ya = year_ago(o)
    rec["year_ago"] = {"date": ya[0], "value": ya[1]} if ya and not no_yoy else None


def _derive_all(cfg, result):
    """Derived series (config 'derived'), KPI tiles ('kpis') and summary table ('summary'),
    using calc.py (same formulas as the Financials Dashboard)."""
    import calc
    S = result["series"]
    get = lambda k: [list(o) for o in (S.get(k, {}).get("obs") or [])]
    for spec in cfg.get("derived", []):
        rec = {k: v for k, v in spec.items()}
        rec.setdefault("fred", "derived")
        ins = calc.inputs_of(spec)
        try:
            obs = calc.derive(spec["op"], spec, get)
        except Exception as e:  # noqa: BLE001
            obs, rec["error"] = [], str(e)
        rec["obs"] = obs
        rec["derived_from"] = ins
        st = [S.get(k, {}).get("status", "failed") for k in ins]
        rec["status"] = "failed" if not obs else ("stale" if any(x != "ok" for x in st) else "ok")
        if obs:
            _latest_fields(rec, spec.get("no_yoy"))
        S[spec["key"]] = rec
        lt = rec.get("latest") or {}
        print(f"  {'= ' + spec['key']:<24} {rec['status']:<7} {lt.get('date','-'):<11} {lt.get('value','')}")
    hs = [h for h, _ in calc.HORIZONS]
    if cfg.get("kpis"):
        out = []
        for k in cfg["kpis"]:
            s = S.get(k["key"]) or {}
            obs = s.get("obs") or []
            t = {**k, "freq": s.get("freq"), "status": s.get("status", "failed"), "src_label": s.get("src_label")}
            if obs:
                t.update(latest=obs[-1][1], asof=obs[-1][0],
                         chg={h: calc.change(obs, k["unit"], h, s.get("freq")) for h in hs})
            out.append(t)
        result["kpis"] = out
    if cfg.get("summary"):
        rows = []
        for section, key, name, unit, pol in cfg["summary"]:
            s = S.get(key) or {}
            obs = s.get("obs") or []
            r = {"section": section, "id": key, "name": name, "unit": unit, "polarity": pol,
                 "change_unit": calc.UNIT_LABEL[unit], "status": s.get("status", "failed")}
            if obs:
                chg = {h: calc.change(obs, unit, h, s.get("freq")) for h in hs}
                r.update(units=s.get("units"), freq=s.get("freq"), latest=obs[-1][1], asof=obs[-1][0],
                         latest_fmt=calc.fmt_level(obs[-1][1], s.get("units")), chg=chg,
                         chg_fmt={h: calc.fmt_change(c, unit) for h, c in chg.items()})
            else:
                r["missing"] = True
            rows.append(r)
        result["summary"] = rows
    if cfg.get("quarterly"):
        result["quarterly"] = _quarterly(cfg["quarterly"], S)
    if cfg.get("growth5y"):
        result["growth5y"] = _growth5y(cfg["growth5y"], S)


def _growth5y(gc, S):
    """Last N quarters (+ QTD) of YoY and compounded LQA growth on two bases: end-of-period (last weekly
    reading on/before quarter-end) and quarter-average."""
    import calc
    per = {r["key"]: calc.quarter_averages((S.get(r["key"]) or {}).get("obs") or [], (S.get(r["key"]) or {}).get("freq")) for r in gc["rows"]}
    anchor = per[gc["anchor"]]
    full = sorted(k for k, v in anchor.items() if v["complete"])
    cols = full[-int(gc.get("quarters", 20)):] + [k for k in sorted(anchor) if k > full[-1]]
    prev_q = lambda y, q: (y, q - 1) if q > 1 else (y - 1, 4)
    g = lambda a, b: round((a / b - 1) * 100, 2) if b else None
    rows = []
    for r in gc["rows"]:
        s = S.get(r["key"]) or {}
        qa, cells = per[r["key"]], []
        for (y, q) in cols:
            c = qa.get((y, q))
            if not c:
                cells.append(None); continue
            yo, lq = qa.get((y - 1, q)), qa.get(prev_q(y, q))
            cell = {"n": c["n"], "partial": not c["complete"], "eop_date": c["last"]}
            for basis, f in (("eop", "eop"), ("avg", "avg")):
                cell[basis] = {"level": round(c[f], 2),
                               "yoy": g(c[f], yo[f]) if yo and yo["complete"] else None,
                               "lqa": round(((c[f] / lq[f]) ** 4 - 1) * 100, 2) if lq and lq["complete"] and lq[f] > 0 and c[f] > 0 else None}
            cells.append(cell)
        rows.append({**r, "sa": r.get("sa", "SA"), "freq": s.get("freq"), "status": s.get("status", "failed"), "cells": cells})
    return {"title": gc.get("title"), "note": gc.get("note"), "rows": rows,
            "columns": [{"label": f"Q{q} {y}", "y": y, "q": q, "eop_date": anchor[(y, q)]["last"], "partial": not anchor[(y, q)]["complete"], "n": anchor[(y, q)]["n"]} for y, q in cols]}


def _quarterly(qc, S):
    """Quarter-average levels with YoY (vs same quarter a year earlier) and compounded LQA
    ((avg_q / avg_q-1)^4 - 1) growth, for the last N quarters. Partial (QTD) quarters are flagged per cell."""
    import calc
    nq = int(qc.get("quarters", 8))
    per = {}
    for g in qc["groups"]:
        for r in g["rows"]:
            s = S.get(r["key"]) or {}
            per[r["key"]] = calc.quarter_averages(s.get("obs") or [], s.get("freq"))
    # columns: the latest nq quarters that are complete for the anchor series, plus any later quarter with data
    anchor = per[qc["anchor"]]
    full = sorted(k for k, v in anchor.items() if v["complete"])
    allq = sorted({k for v in per.values() for k in v})
    cols = full[-nq:] + [k for k in allq if k > full[-1]]
    prev_q = lambda y, q: (y, q - 1) if q > 1 else (y - 1, 4)
    groups = []
    for g in qc["groups"]:
        rows = []
        for r in g["rows"]:
            s = S.get(r["key"]) or {}
            qa = per[r["key"]]
            cells = []
            for (y, q) in cols:
                c = qa.get((y, q))
                if not c:
                    cells.append(None); continue
                yo, lq = qa.get((y - 1, q)), qa.get(prev_q(y, q))
                cells.append({"level": round(c["avg"], 2), "n": c["n"], "partial": not c["complete"], "last": c["last"],
                              "yoy": round((c["avg"] / yo["avg"] - 1) * 100, 2) if yo and yo["complete"] and yo["avg"] else None,
                              "lqa": round(((c["avg"] / lq["avg"]) ** 4 - 1) * 100, 2) if lq and lq["complete"] and lq["avg"] > 0 and c["avg"] > 0 else None})
            rows.append({**r, "freq": s.get("freq"), "status": s.get("status", "failed"), "cells": cells})
        groups.append({"title": g["title"], "bank": g.get("bank"), "rows": rows})
    return {"columns": [{"label": f"Q{q} {y}", "y": y, "q": q} for y, q in cols], "groups": groups,
            "note": qc.get("note"), "title": qc.get("title")}


def quarterly_table(qd):
    """Flat rows (header + data) with every column: level, YoY and LQA per quarter."""
    hdr = ["Group", "Series", "FRED id", "SA/NSA", "Frequency"]
    for c in qd["columns"]:
        hdr += [f"{c['label']} avg level ($bn)", f"{c['label']} YoY %", f"{c['label']} LQA %", f"{c['label']} obs in qtr", f"{c['label']} partial (QTD)"]
    out = [hdr]
    for g in qd["groups"]:
        for r in g["rows"]:
            row = [g["title"], r["name"], r["key"], r.get("sa", "SA"), r.get("freq")]
            for c in r["cells"]:
                row += [None] * 5 if not c else [c["level"], c["yoy"], c["lqa"], c["n"], "QTD" if c["partial"] else ""]
            out.append(row)
    return out


def _write_quarterly_files(did, qd):
    rows = quarterly_table(qd)
    with open(DATA / f"{did}-quarterly.csv", "w", newline="") as f:
        csv.writer(f).writerows([["" if v is None else v for v in r] for r in rows])
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        print("  (openpyxl missing: skipped xlsx)"); return
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Quarterly"
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "F2"
    ws.append([]); ws.append([qd.get("note") or ""])
    wb.properties.created = wb.properties.modified = dt.datetime(2000, 1, 1)  # deterministic file -> no spurious commits
    wb.save(DATA / f"{did}-quarterly.xlsx")


def growth5y_table(gd, basis):
    hdr = ["Bank group", "Basis", "Series", "FRED id", "SA/NSA"]
    BG = {"D": "Domestic total", "L": "Large domestic", "S": "Small domestic"}
    for c in gd["columns"]:
        lab = c["label"] + (" QTD" if c["partial"] else "")
        hdr += [f"{lab} level ($bn)", f"{lab} YoY %", f"{lab} LQA %"] + ([f"{lab} EOP week"] if basis == "eop" else [f"{lab} weeks in qtr"])
    out = [hdr]
    for r in gd["rows"]:
        row = [BG.get(r.get("bank"), "All commercial banks"), "End of period" if basis == "eop" else "Quarter average", "  " * r.get("lvl", 0) + r["name"], r["key"], r["sa"]]
        for c in r["cells"]:
            if not c:
                row += [None] * 4; continue
            b = c[basis]
            row += [b["level"], b["yoy"], b["lqa"], c["eop_date"] if basis == "eop" else c["n"]]
        out.append(row)
    return out


def _write_growth5y_files(did, gd):
    eop, avg = growth5y_table(gd, "eop"), growth5y_table(gd, "avg")
    with open(DATA / f"{did}-growth5y.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerows([["" if v is None else v for v in r] for r in eop])
        w.writerow([])
        w.writerows([["" if v is None else v for v in r] for r in avg])
        w.writerows([[], [gd.get("note") or ""]])
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        print("  (openpyxl missing: skipped xlsx)"); return
    wb = openpyxl.Workbook()
    for i, (name, rows) in enumerate((("EOP", eop), ("Average", avg))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "F2"
        ws.append([]); ws.append([gd.get("note") or ""])
    wb.properties.created = wb.properties.modified = dt.datetime(2000, 1, 1)
    wb.save(DATA / f"{did}-growth5y.xlsx")


def _strip(doc):
    """Comparable view of a dashboard doc, without run timestamps."""
    doc = json.loads(json.dumps(doc))  # normalize tuples -> lists
    return {"config": doc.get("config"), "kpis": doc.get("kpis"), "summary": doc.get("summary"), "auctions": doc.get("auctions"), "deposit_pricing": doc.get("deposit_pricing"),
            "series": {k: {f: v for f, v in s.items() if f != "fetched_at"} for k, s in doc.get("series", {}).items()}}


def main():
    hub = json.loads((ROOT / "config/hub.json").read_text())
    only = set(sys.argv[1:])
    bad = 0
    status_path = DATA / "status.json"
    status = json.loads(status_path.read_text()) if status_path.exists() else {}
    status.setdefault("dashboards", {})
    for a in hub["areas"]:
        if a.get("status") != "live" or not a.get("dashboard"):
            continue
        if only and a["id"] not in only:
            continue
        print(f"[{a['id']}] {a['dashboard']}")
        r = run_dashboard(a["dashboard"])
        bad += sum(1 for s in r["series"].values() if s["status"] != "ok")
        status["dashboards"][a["id"]] = {
            "checked_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
            "data_changed_at": r["updated_at"],
            "series": {k: {"fred": s.get("fred"), "status": s["status"], "latest": s.get("latest"), **({"error": s["error"]} if s.get("error") else {})}
                       for k, s in r["series"].items()}}
    status["checked_at"] = dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
    status["ok"] = bad == 0
    status_path.write_text(json.dumps(status, indent=1))
    (DATA / "hub.json").write_text(json.dumps(hub, indent=1))
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
